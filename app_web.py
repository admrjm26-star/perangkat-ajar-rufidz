from io import BytesIO
import os
import time
from typing import List
import warnings

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from google import genai
from google.genai import types
from pydantic import BaseModel, Field
import streamlit as st

# Sembunyikan warning
warnings.filterwarnings("ignore")

# ---------------------------------------------------------
# 1. KONFIGURASI HALAMAN STREAMLIT
# ---------------------------------------------------------
st.set_page_config(
    page_title="Perangkat Ajar The Rufidz Indonesia",
    page_icon="📖",
    layout="wide",
)

# Custom CSS Tampilan Pesantren & Menyembunyikan Header/GitHub
st.markdown(
    """
    <style>
    /* Sembunyikan Header, Menu, Footer & Toolbar Streamlit Cloud */
    header {visibility: hidden;}
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    [data-testid="stHeader"] {display: none;}
    [data-testid="stToolbar"] {display: none;}

    .main-title {
        color: #0A5C36;
        text-align: center;
        font-weight: bold;
        margin-bottom: 0px;
    }
    .sub-title {
        text-align: center;
        color: #555555;
        font-style: italic;
        margin-bottom: 30px;
    }
    .stButton>button {
        background-color: #0A5C36;
        color: white;
        font-weight: bold;
        border-radius: 8px;
        height: 3em;
        width: 100%;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# API KEY GEMINI & INISIALISASI CLIENT
client = genai.Client(api_key=st.secrets["GEMINI_API_KEY"])


# ---------------------------------------------------------
# 2. SCHEMA PYDANTIC DATA (DENGAN FIELD SILABUS DIPERBARUI)
# ---------------------------------------------------------
class SilabusItem(BaseModel):
    pertemuan_ke: int
    bab_fasal: str
    perkiraan_halaman: str
    alokasi_waktu: str
    metode: str
    capaian_indikator: str
    indikator_ketercapaian: str
    bentuk_evaluasi: str


class RPPKegiatanInti(BaseModel):
    kegiatan_literasi: str
    critical_thinking: str
    collaboration: str
    communication: str
    creativity: str


class RPPDetail(BaseModel):
    sekolah: str = "The Rufidz Indonesia"
    mata_pelajaran: str
    kelas_tahap: str
    materi_pokok: str
    alokasi_waktu: str
    tujuan_pembelajaran: List[str]
    pendahuluan: List[str]
    kegiatan_inti: RPPKegiatanInti
    penutup: List[str]
    penilaian: List[str]


class PerangkatAjarRufidz(BaseModel):
    nama_kitab: str
    fan_ilmu: str
    tingkat_rekomendasi: str
    total_pertemuan: int
    silabus: List[SilabusItem]
    sample_rpp_pertemuan_1: RPPDetail


# ---------------------------------------------------------
# 3. FUNGSI GENERATE AI (GEMINI)
# ---------------------------------------------------------
def generate_ai(nama_kitab: str, tingkat_kelas: str) -> PerangkatAjarRufidz:
    system_prompt = (
        "Kamu adalah pakar kurikulum lembaga Rufidz Tahfidz & Diniyah Indonesia. "
        "Tugasmu menyusun Perangkat Ajar (Silabus & RPP 1 Lembar) berbasis nama kitab/pelajaran. "
        "Setiap item Silabus HARUS melengkapi: pertemuan_ke, bab_fasal, perkiraan_halaman, alokasi_waktu, "
        "metode (klasik), capaian_indikator, indikator_ketercapaian, dan bentuk_evaluasi. "
        "RPP mencakup 5 komponen kegiatan inti: Kegiatan Literasi, Critical Thinking, "
        "Collaboration, Communication, dan Creativity. "
        "Gunakan METODE PEMBELAJARAN KLASIK / KLASIKAL yang umum (seperti Ceramah Interaktif, "
        "Tanya Jawab, Diskusi, Demonstrasi, Drill/Latihan, Muroja'ah, dan Penugasan). "
        "Jangan menggunakan istilah Sorogan atau Bandongan. "
        f"PENTING: Tingkat/Kelas HARUS disesuaikan secara presisi dengan pilihan user: '{tingkat_kelas}'. "
        "Pastikan output STRICT mengikuti skema JSON."
    )

    max_retries = 5
    for attempt in range(1, max_retries + 1):
        try:
            response = client.models.generate_content(
                model="gemini-3.6-flash",
                contents=f"Buatkan silabus dan RPP 1 Lembar lengkap untuk kitab/materi: {nama_kitab} khusus untuk tingkat/kelas: {tingkat_kelas}",
                config=types.GenerateContentConfig(
                    system_instruction=system_prompt,
                    response_mime_type="application/json",
                    response_schema=PerangkatAjarRufidz,
                    temperature=0.2,
                ),
            )
            return PerangkatAjarRufidz.model_validate_json(response.text)
        except Exception as e:
            if "503" in str(e) or "UNAVAILABLE" in str(e):
                time.sleep(4)
            else:
                raise e
    raise Exception("Gagal terhubung ke server AI. Silakan coba kembali.")


# ---------------------------------------------------------
# 4. FUNGSI EKSPOR KE WORD (.DOCX)
# ---------------------------------------------------------

# A. EKSPOR RPP SAJA
def create_rpp_docx_bytes(data: PerangkatAjarRufidz) -> BytesIO:
    doc = docx.Document()
    PRIMARY_RGB = RGBColor(10, 92, 54)

    p_kop = doc.add_paragraph()
    p_kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kop.paragraph_format.space_after = Pt(2)

    r_brand = p_kop.add_run("THE RUFIDZ INDONESIA\n")
    r_brand.font.size = Pt(14)
    r_brand.font.bold = True
    r_brand.font.color.rgb = PRIMARY_RGB

    r_title = p_kop.add_run("RENCANA PELAKSANAAN PEMBELAJARAN (RPP)")
    r_title.font.size = Pt(12)
    r_title.font.bold = True

    rpp = data.sample_rpp_pertemuan_1

    p_rpp_id = doc.add_paragraph()
    p_rpp_id.paragraph_format.space_after = Pt(6)
    p_rpp_id.add_run(f"Sekolah       : {rpp.sekolah}\n")
    p_rpp_id.add_run(f"Mata Pelajaran: {rpp.mata_pelajaran}\n")
    p_rpp_id.add_run(f"Kelas / Tahap : {rpp.kelas_tahap}\n")
    p_rpp_id.add_run(f"Materi Pokok  : {rpp.materi_pokok}\n")
    p_rpp_id.add_run(f"Alokasi Waktu : {rpp.alokasi_waktu}")

    doc.add_heading("A. Tujuan Pembelajaran", level=3)
    for t in rpp.tujuan_pembelajaran:
        doc.add_paragraph(f"• {t}")

    doc.add_heading("B. Langkah-Langkah Pembelajaran", level=3)
    doc.add_heading("1. Kegiatan Pendahuluan", level=4)
    for p_step in rpp.pendahuluan:
        doc.add_paragraph(f"• {p_step}")

    doc.add_heading("2. Kegiatan Inti", level=4)
    t_inti = doc.add_table(rows=1, cols=2)
    t_inti.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_inti.rows[0].cells[0].text = "Aspek"
    t_inti.rows[0].cells[1].text = "Kegiatan Pembelajaran"
    for cell in t_inti.rows[0].cells:
        cell.paragraphs[0].runs[0].font.bold = True

    aspek_list = [
        ("Kegiatan Literasi", rpp.kegiatan_inti.kegiatan_literasi),
        ("Critical Thinking", rpp.kegiatan_inti.critical_thinking),
        ("Collaboration", rpp.kegiatan_inti.collaboration),
        ("Communication", rpp.kegiatan_inti.communication),
        ("Creativity", rpp.kegiatan_inti.creativity),
    ]
    for asp, desk in aspek_list:
        row = t_inti.add_row().cells
        row[0].text = asp
        row[1].text = desk

    doc.add_heading("3. Kegiatan Penutup", level=4)
    for p_close in rpp.penutup:
        doc.add_paragraph(f"• {p_close}")

    doc.add_heading("C. Penilaian Hasil Pembelajaran", level=3)
    for pen in rpp.penilaian:
        doc.add_paragraph(f"• {pen}")

    doc.add_paragraph("\n")
    p_ttd = doc.add_paragraph()
    p_ttd.add_run("Mengetahui,\n")
    p_ttd.add_run("Kepala Sekolah\t\t\t\tPengampu Mapel\n\n\n\n")
    p_ttd.add_run("( Mudir Sekolah )\t\t\t( Ustadz Pengampu )")

    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


# B. EKSPOR SILABUS SAJA (8 KOLOM)
def create_silabus_docx_bytes(data: PerangkatAjarRufidz) -> BytesIO:
    doc = docx.Document()
    PRIMARY_RGB = RGBColor(10, 92, 54)

    p_kop = doc.add_paragraph()
    p_kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kop.paragraph_format.space_after = Pt(2)

    r_brand = p_kop.add_run("THE RUFIDZ INDONESIA\n")
    r_brand.font.size = Pt(14)
    r_brand.font.bold = True
    r_brand.font.color.rgb = PRIMARY_RGB

    r_title = p_kop.add_run("SILABUS PEMBELAJARAN")
    r_title.font.size = Pt(12)
    r_title.font.bold = True

    doc.add_heading("I. IDENTITAS MATA PELAJARAN", level=2)
    p_id = doc.add_paragraph()
    p_id.add_run(f"Nama Kitab/Pelajaran: {data.nama_kitab}\n").bold = True
    p_id.add_run(f"Fan Ilmu: {data.fan_ilmu}\n")
    p_id.add_run(f"Tingkat / Kelas: {data.tingkat_rekomendasi}\n")
    p_id.add_run(f"Total Pertemuan: {data.total_pertemuan} Pertemuan")

    doc.add_heading("II. SILABUS PEMBELAJARAN", level=2)
    table = doc.add_table(rows=1, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    headers = [
        "Ptm",
        "Bab / Fasal",
        "Halaman",
        "Alokasi Waktu",
        "Metode Klasik",
        "Capaian Indikator",
        "Indikator Ketercapaian",
        "Bentuk Evaluasi",
    ]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True

    for item in data.silabus:
        row_cells = table.add_row().cells
        row_cells[0].text = str(item.pertemuan_ke)
        row_cells[1].text = item.bab_fasal
        row_cells[2].text = item.perkiraan_halaman
        row_cells[3].text = item.alokasi_waktu
        row_cells[4].text = item.metode
        row_cells[5].text = item.capaian_indikator
        row_cells[6].text = item.indikator_ketercapaian
        row_cells[7].text = item.bentuk_evaluasi

    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


# C. EKSPOR PAKET LENGKAP (SILABUS 8 KOLOM + RPP)
def create_full_docx_bytes(data: PerangkatAjarRufidz) -> BytesIO:
    doc = docx.Document()
    PRIMARY_RGB = RGBColor(10, 92, 54)

    # 1. Silabus Section
    p_kop = doc.add_paragraph()
    p_kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_brand = p_kop.add_run("PERANGKAT AJAR THE RUFIDZ INDONESIA\n")
    r_brand.font.size = Pt(16)
    r_brand.font.bold = True
    r_brand.font.color.rgb = PRIMARY_RGB

    r_title = p_kop.add_run("SILABUS & RENCANA PELAKSANAAN PEMBELAJARAN (RPP)")
    r_title.font.size = Pt(12)
    r_title.font.bold = True

    doc.add_heading("I. IDENTITAS PERANGKAT AJAR", level=2)
    p_id = doc.add_paragraph()
    p_id.add_run(f"Nama Kitab/Pelajaran: {data.nama_kitab}\n").bold = True
    p_id.add_run(f"Fan Ilmu: {data.fan_ilmu}\n")
    p_id.add_run(f"Tingkat / Kelas: {data.tingkat_rekomendasi}\n")
    p_id.add_run(f"Total Pertemuan: {data.total_pertemuan} Pertemuan")

    doc.add_heading("II. SILABUS PEMBELAJARAN", level=2)
    table = doc.add_table(rows=1, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    headers = [
        "Ptm",
        "Bab / Fasal",
        "Halaman",
        "Alokasi Waktu",
        "Metode Klasik",
        "Capaian Indikator",
        "Indikator Ketercapaian",
        "Bentuk Evaluasi",
    ]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True

    for item in data.silabus:
        row_cells = table.add_row().cells
        row_cells[0].text = str(item.pertemuan_ke)
        row_cells[1].text = item.bab_fasal
        row_cells[2].text = item.perkiraan_halaman
        row_cells[3].text = item.alokasi_waktu
        row_cells[4].text = item.metode
        row_cells[5].text = item.capaian_indikator
        row_cells[6].text = item.indikator_ketercapaian
        row_cells[7].text = item.bentuk_evaluasi

    # PAGE BREAK SUPAYA RPP MULAI DI HALAMAN BARU
    doc.add_page_break()

    # 2. RPP Section
    doc.add_heading("III. RENCANA PELAKSANAAN PEMBELAJARAN (RPP)", level=2)
    rpp = data.sample_rpp_pertemuan_1

    p_rpp_id = doc.add_paragraph()
    p_rpp_id.add_run(f"Sekolah       : {rpp.sekolah}\n")
    p_rpp_id.add_run(f"Mata Pelajaran: {rpp.mata_pelajaran}\n")
    p_rpp_id.add_run(f"Kelas / Tahap : {rpp.kelas_tahap}\n")
    p_rpp_id.add_run(f"Materi Pokok  : {rpp.materi_pokok}\n")
    p_rpp_id.add_run(f"Alokasi Waktu : {rpp.alokasi_waktu}")

    doc.add_heading("A. Tujuan Pembelajaran", level=3)
    for t in rpp.tujuan_pembelajaran:
        doc.add_paragraph(f"• {t}")

    doc.add_heading("B. Langkah-Langkah Pembelajaran", level=3)
    doc.add_heading("Kegiatan Pendahuluan", level=4)
    for p_step in rpp.pendahuluan:
        doc.add_paragraph(f"• {p_step}")

    doc.add_heading("Kegiatan Inti", level=4)
    t_inti = doc.add_table(rows=1, cols=2)
    t_inti.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_inti.rows[0].cells[0].text = "Aspek"
    t_inti.rows[0].cells[1].text = "Kegiatan Pembelajaran"

    aspek_list = [
        ("Kegiatan Literasi", rpp.kegiatan_inti.kegiatan_literasi),
        ("Critical Thinking", rpp.kegiatan_inti.critical_thinking),
        ("Collaboration", rpp.kegiatan_inti.collaboration),
        ("Communication", rpp.kegiatan_inti.communication),
        ("Creativity", rpp.kegiatan_inti.creativity),
    ]
    for asp, desk in aspek_list:
        row = t_inti.add_row().cells
        row[0].text = asp
        row[1].text = desk

    doc.add_heading("Kegiatan Penutup", level=4)
    for p_close in rpp.penutup:
        doc.add_paragraph(f"• {p_close}")

    doc.add_heading("C. Penilaian Hasil Pembelajaran", level=3)
    for pen in rpp.penilaian:
        doc.add_paragraph(f"• {pen}")

    doc.add_paragraph("\n")
    p_ttd = doc.add_paragraph()
    p_ttd.add_run("Mengetahui,\n")
    p_ttd.add_run("Kepala Sekolah\t\t\t\tPengampu Mapel\n\n\n\n")
    p_ttd.add_run("( Mudir Sekolah )\t\t\t( Ustadz Pengampu )")

    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


# ---------------------------------------------------------
# 5. TAMPILAN UTAMA APLIKASI WEB
# ---------------------------------------------------------
st.markdown(
    "<h1 class='main-title'>PERANGKAT AJAR THE RUFIDZ INDONESIA</h1>",
    unsafe_allow_html=True,
)
st.markdown(
    "<p class='sub-title'>Generator Silabus & RPP Otomatis untuk Rufidz -Tahfidz & Diniyah</p>",
    unsafe_allow_html=True,
)

# Sidebar
st.sidebar.header("📌 Referensi Cepat Kitab")
preset_kitab = st.sidebar.selectbox(
    "Pilih Kitab Populer:",
    [
        "-- Pilih atau Ketik Manual --",
        "Al-Ushul Ats-Tsalatsah",
        "Matan Al-Jurumiyah",
        "Tuhfah Saniyyah",
        "Al-Qawa’idun Nahwiyyah",
        "Qawa’id Arba’",
        "Laamiyyah",
        "Ushulun Fit Tafsir",
        "Al-Khuluqul Hasan",
    ],
)

# Input Nama Kitab & Pilih Tingkat/Kelas
col_input1, col_input2 = st.columns([2, 1])

with col_input1:
    default_val = (
        "" if preset_kitab == "-- Pilih atau Ketik Manual --" else preset_kitab
    )
    nama_kitab_input = st.text_input(
        "Nama Kitab / Pelajaran:",
        value=default_val,
        placeholder="Contoh: Qawa’id Arba’ / Jurumiyah",
    )

with col_input2:
    tingkat_kelas_input = st.selectbox(
        "Pilih Tingkat / Kelas:",
        ["Kelas A", "Kelas B", "Kelas C", "Kelas 7", "Kelas 8", "Kelas 9", "KPM Pro"],
    )

# Tombol Generate
if st.button("🚀 Generate Silabus & RPP Otomatis"):
    if not nama_kitab_input.strip():
        st.warning("Silakan masukkan atau pilih nama kitab/pelajaran terlebih dahulu!")
    else:
        with st.spinner(
            f"Sedang menyusun kurikulum '{nama_kitab_input}' untuk {tingkat_kelas_input}..."
        ):
            try:
                # 1. AI Generate
                result = generate_ai(nama_kitab_input, tingkat_kelas_input)
                st.session_state["data_result"] = result
                st.success("Berhasil di-generate oleh AI!")
            except Exception as e:
                st.error(f"Terjadi kesalahan: {e}")

# Tampilkan Hasil Jika Ada
if "data_result" in st.session_state:
    data = st.session_state["data_result"]

    st.divider()
    st.subheader(f"📖 Identitas Kurikulum: {data.nama_kitab}")

    col1, col2, col3 = st.columns(3)
    col1.metric("Fan Ilmu", data.fan_ilmu)
    col2.metric("Tingkat / Kelas", data.tingkat_rekomendasi)
    col3.metric("Total Pertemuan", f"{data.total_pertemuan} Pertemuan")

    tab1, tab2 = st.tabs(["📋 Silabus Pembelajaran", "📝 Sample RPP 1 Lembar"])

    with tab1:
        # Tampilkan Tabel Silabus (8 Kolom)
        table_data = []
        for s in data.silabus:
            table_data.append({
                "Pertemuan": s.pertemuan_ke,
                "Bab / Fasal": s.bab_fasal,
                "Perkiraan Hal.": s.perkiraan_halaman,
                "Alokasi Waktu": s.alokasi_waktu,
                "Metode Klasik": s.metode,
                "Capaian Indikator": s.capaian_indikator,
                "Indikator Ketercapaian": s.indikator_ketercapaian,
                "Bentuk Evaluasi": s.bentuk_evaluasi,
            })
        st.dataframe(table_data, use_container_width=True)

    with tab2:
        rpp = data.sample_rpp_pertemuan_1
        st.markdown("### RENCANA PELAKSANAAN PEMBELAJARAN (RPP)")
        st.text(f"Sekolah       : {rpp.sekolah}")
        st.text(f"Mata Pelajaran: {rpp.mata_pelajaran}")
        st.text(f"Kelas / Tahap : {rpp.kelas_tahap}")
        st.text(f"Materi Pokok  : {rpp.materi_pokok}")
        st.text(f"Alokasi Waktu : {rpp.alokasi_waktu}")

        st.markdown("#### A. Tujuan Pembelajaran")
        for t in rpp.tujuan_pembelajaran:
            st.markdown(f"- {t}")

        st.markdown("#### B. Langkah-Langkah Pembelajaran")
        st.markdown("**1. Kegiatan Pendahuluan:**")
        for p_step in rpp.pendahuluan:
            st.markdown(f"- {p_step}")

        st.markdown("**2. Kegiatan Inti:**")
        st.table([
            {"Aspek": "Kegiatan Literasi", "Deskripsi": rpp.kegiatan_inti.kegiatan_literasi},
            {"Aspek": "Critical Thinking", "Deskripsi": rpp.kegiatan_inti.critical_thinking},
            {"Aspek": "Collaboration", "Deskripsi": rpp.kegiatan_inti.collaboration},
            {"Aspek": "Communication", "Deskripsi": rpp.kegiatan_inti.communication},
            {"Aspek": "Creativity", "Deskripsi": rpp.kegiatan_inti.creativity},
        ])

        st.markdown("**3. Kegiatan Penutup:**")
        for p_close in rpp.penutup:
            st.markdown(f"- {p_close}")

        st.markdown("#### C. Penilaian Hasil Pembelajaran")
        for pen in rpp.penilaian:
            st.markdown(f"- {pen}")

    st.divider()

    # OPSI DOWNLOAD TERPISAH
    st.subheader("📥 Pilih File yang Ingin Di-download:")
    col_dl1, col_dl2, col_dl3 = st.columns(3)

    clean_name = "".join(
        c for c in data.nama_kitab if c.isalnum() or c in (" ", "_")
    ).rstrip()

    with col_dl1:
        st.download_button(
            label="📄 Download RPP Saja (.docx)",
            data=create_rpp_docx_bytes(data),
            file_name=f"RPP_{clean_name.replace(' ', '_')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    with col_dl2:
        st.download_button(
            label="📋 Download Silabus Saja (.docx)",
            data=create_silabus_docx_bytes(data),
            file_name=f"Silabus_{clean_name.replace(' ', '_')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    with col_dl3:
        st.download_button(
            label="📦 Download Paket Lengkap (.docx)",
            data=create_full_docx_bytes(data),
            file_name=f"Perangkat_Ajar_Lengkap_{clean_name.replace(' ', '_')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
